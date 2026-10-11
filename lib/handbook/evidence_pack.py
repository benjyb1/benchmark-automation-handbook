"""Evidence on request: for chosen rows (and, if asked, chosen questions), every answer with the
words it rests on, as an Excel workbook (one sheet per row) or one readable HTML page that can be
sent to the organisation scored.

What goes in: the cell's state, answer and points, each quote (with its English translation), the
document's title, publisher, link, page and date, how strong the source is, the cell's flags and
the rules it rests on (the rule's question, or its decision). What never goes in: anything from
the held-back folder or `compare/`, a rule's note about last year or its reasoning, and the
scorer's `note` and `unsure`. This module reads only `project`, `questions`, `units`, `rules`,
`answers/`, `checks/` and document metas, and picks named keys from each, so keep it that way.

The HTML is fixed copy: the words around the data are written here, never generated.
"""
from __future__ import annotations
import html
import os
import re
from datetime import datetime, timezone
from pathlib import Path

from .bundle import FLAG_WORDS, STATE_WORDS
from .folder import HandbookError, Project, applies, effective_question

FORMATS = ("xlsx", "html")
COLUMNS = ("Question", "Wording", "Answer", "What it means", "Points", "Out of", "State", "Quote",
           "Quote in English", "Document title", "Publisher", "Link", "Page", "Document date",
           "Source strength", "Flags", "Rules it rests on")

OWN_WORDS = "Own words"
THIRD_PARTY = "published by someone other than the organisation"
NO_REASON = "no reason recorded"
NO_EVIDENCE = "No evidence was found for this question."
NOT_RUN = "This question has not been scored yet."
NOT_APPLICABLE = "This question does not apply to this {word}."
NO_QUOTE = "No quote was recorded for this answer."
MISSING_DOCUMENT = "This document is no longer in the project."
OUT_OF_DATE = "Out of date: to be scored again"

_MONTHS = ("January", "February", "March", "April", "May", "June", "July", "August", "September",
           "October", "November", "December")


# --- choosing rows and questions ------------------------------------------------------------

def _words(project: Project) -> tuple[str, str]:
    proj = project.read("project")
    return str(proj.get("unit_word") or "row"), str(proj.get("unit_word_plural") or "rows")


def _listed(values, limit=12) -> str:
    values = [str(v) for v in values]
    more = f" and {len(values) - limit} more" if len(values) > limit else ""
    return ", ".join(values[:limit]) + more


def choose_rows(project: Project, rows: list[str] | None = None, label: str | None = None,
                all_rows: bool = False) -> list[dict]:
    """The units chosen by id (`rows`), by one label value (`label` as "LABEL=VALUE", label and
    value matched ignoring case) or `all_rows`, in units.json order. Exactly one way must be
    given; raises HandbookError for an unknown id, label or value."""
    if sum((rows is not None, label is not None, bool(all_rows))) != 1:
        raise HandbookError("Choose the rows one way: by id (--rows), by a label value "
                            "(--label LABEL=VALUE) or all of them (--all).")
    word, words = _words(project)
    units = [u for u in project.read("units").get("units") or [] if isinstance(u, dict)]
    if not units:
        raise HandbookError(f"There are no {words} in units.json yet, so there is no evidence to export.")
    if all_rows:
        return units
    if rows is not None:
        wanted = [r for r in dict.fromkeys(rows) if r]
        if not wanted:
            raise HandbookError(f"Give at least one {word} id with --rows.")
        known = {u.get("id") for u in units}
        unknown = [r for r in wanted if r not in known]
        if unknown:
            raise HandbookError(f"There is no {word} with the id {_listed(unknown)} in units.json. "
                                f"The ids are: {_listed(u.get('id') for u in units)}.")
        return [u for u in units if u.get("id") in wanted]
    name, sep, value = str(label).partition("=")
    name, value = name.strip(), value.strip()
    if not sep or not name or not value:
        raise HandbookError("Give the label as LABEL=VALUE, for example country=Denmark.")
    labels = list(dict.fromkeys([*(project.read("project").get("labels") or []),
                                 *(k for u in units for k in (u.get("labels") or {}))]))
    match = next((lb for lb in labels if lb == name), None) or \
        next((lb for lb in labels if lb.casefold() == name.casefold()), None)
    if match is None:
        raise HandbookError(f"There is no label {name!r}. The labels are: {_listed(labels) or 'none'}.")
    values = [(u.get("labels") or {}).get(match) for u in units]
    chosen = [u for u, v in zip(units, values) if str(v).strip().casefold() == value.casefold()]
    if not chosen:
        there = sorted({str(v) for v in values if v not in (None, "")})
        raise HandbookError(f"No {words} have {match} = {value}. The values of {match} are: "
                            f"{_listed(there) or 'none'}.")
    return chosen


def choose_questions(project: Project, questions: list[str] | None = None) -> list[str] | None:
    """The question ids asked for, in questions.json order; None (all) when none were given.
    Raises HandbookError naming any id that is not in questions.json."""
    if not questions:
        return None
    known = [q.get("id") for q in project.read("questions").get("questions") or []]
    unknown = [q for q in dict.fromkeys(questions) if q not in known]
    if unknown:
        raise HandbookError(f"There is no question {_listed(unknown)} in questions.json. "
                            f"The questions are: {_listed(known)}.")
    return [q for q in known if q in questions]


# --- what a line says -----------------------------------------------------------------------

def source_strength(meta: dict | None) -> str | None:
    """"Own words", or "Weak: why" when the document's meta has `weak: {"why"}` or it is a
    third-party copy the owner did not provide. Owner-provided documents are never weak. None
    when there is no meta (the document is no longer in the project)."""
    if meta is None:
        return None
    if meta.get("provided_by_owner"):
        return OWN_WORDS
    weak = meta.get("weak")
    if weak:
        why = str(weak.get("why") or "").strip() if isinstance(weak, dict) else ""
        return f"Weak: {why or NO_REASON}"
    if meta.get("host_type") == "third_party":
        return f"Weak: {THIRD_PARTY}"
    return OWN_WORDS


def _flag_word(flag: str) -> str:
    return FLAG_WORDS.get(flag) or str(flag).replace("_", " ").capitalize()


def _rule_text(rule_id: str, rules: dict) -> str:
    """The rule's decision when decided, else its question; never its reasoning or last-year note."""
    rule = rules.get(rule_id)
    if not isinstance(rule, dict):
        return str(rule_id)
    if rule.get("status") == "decided":
        decision = rule.get("decision") or {}
        text = decision.get("text") or next((o.get("text") for o in rule.get("options") or []
                                             if o.get("id") == decision.get("option")), None)
        return f"{rule_id}: {text}" if text else str(rule_id)
    return f"{rule_id} (not decided yet): {rule.get('ask') or ''}".rstrip(": ")


def _level(question: dict, value) -> tuple[str | None, int | float | None]:
    """(the meaning of the answer's level, the most points the question can give)."""
    levels = [lv for lv in (question.get("answer") or {}).get("levels") or [] if isinstance(lv, dict)]
    points = [lv.get("points") for lv in levels if isinstance(lv.get("points"), (int, float))]
    level = next((lv for lv in levels if lv.get("value") == value), None) if value is not None else None
    if level is None and value is not None:
        level = next((lv for lv in levels if str(lv.get("value")) == str(value)), None)
    return (level or {}).get("meaning"), (max(points) if points else None)


def _evidence(item: dict, documents: dict) -> dict:
    meta = documents.get(item.get("doc"))
    quote, quote_en = item.get("quote"), item.get("quote_en")
    out = {"doc": item.get("doc"), "quote": quote,
           "quote_en": quote_en if quote_en and quote_en != quote else None,
           "language": item.get("language"), "page": item.get("page")}
    if meta is None:
        out.update(title=MISSING_DOCUMENT, publisher=None, link=None, date=None, strength=None)
    else:
        out.update(title=meta.get("title") or item.get("doc"), publisher=meta.get("publisher"),
                   link=meta.get("url") or meta.get("archive_url"), date=meta.get("date"),
                   strength=source_strength(meta))
    return out


def _entry(question: dict, answer: dict | None, check: dict | None, documents: dict, rules: dict,
           applicable: bool, word: str) -> dict:
    entry = {"question": question.get("id"), "short": question.get("short"),
             "wording": question.get("wording"), "answer": None, "meaning": None, "points": None,
             "out_of": _level(question, None)[1], "evidence": [], "note": None, "flags": [],
             "rules": []}
    if not applicable:
        return {**entry, "state": STATE_WORDS["not_applicable"],
                "note": NOT_APPLICABLE.format(word=word)}
    state = (answer or {}).get("state") or "not_run"
    entry["state"] = STATE_WORDS.get(state, state)
    if answer is None or state == "not_run":
        entry["note"] = NOT_RUN
    else:
        entry["answer"], entry["points"] = answer.get("value"), answer.get("points")
        entry["meaning"] = _level(question, entry["answer"])[0]
        entry["evidence"] = [_evidence(e, documents) for e in answer.get("evidence") or []
                             if isinstance(e, dict)]
        entry["rules"] = [_rule_text(r, rules) for r in answer.get("rules") or []]
        if not entry["evidence"]:
            entry["note"] = NO_EVIDENCE if state == "no_evidence" else NO_QUOTE
    entry["flags"] = [_flag_word(f) for f in (check or {}).get("flags") or []]
    if (answer or {}).get("out_of_date"):
        entry["flags"].append(OUT_OF_DATE)
    return entry


def build_pack(project: Project, units: list[dict], questions: list[str] | None = None,
               now: datetime | None = None) -> dict:
    """{"name", "edition", "unit_word", "unit_word_plural", "made" (ISO date), "questions",
    "rows": [{"row", "name", "labels": [[label, value]], "questions": [entry]}]}. An entry is
    {"question", "short", "wording", "state" (in words), "answer", "meaning", "points", "out_of",
    "evidence": [{"doc", "quote", "quote_en", "language", "page", "title", "publisher", "link",
    "date", "strength"}], "note" (fixed words when there is no quote), "flags" (in words),
    "rules" (in words)}. Every question that applies to a row gets an entry; with `questions`,
    each listed question does, saying so when it does not apply."""
    proj = project.read("project")
    word, words = _words(project)
    labels = list(proj.get("labels") or [])
    documents = project.documents()
    rules = {r.get("id"): r for r in project.read("rules").get("rules") or [] if isinstance(r, dict)}
    all_questions = project.read("questions").get("questions") or []
    made = (now or datetime.now(timezone.utc)).date().isoformat()
    rows = []
    for unit in units:
        row = unit.get("id")
        answers = project.answers(row).get("cells") or {}
        checks = project.checks(row).get("cells") or {}
        entries = []
        for q in all_questions:
            qid = q.get("id")
            if questions is not None and qid not in questions:
                continue
            applicable = applies(q, unit)
            if not applicable and questions is None:
                continue
            entries.append(_entry(effective_question(q, unit), answers.get(qid), checks.get(qid),
                                  documents, rules, applicable, word))
        unit_labels = unit.get("labels") or {}
        order = list(dict.fromkeys([*labels, *unit_labels]))
        rows.append({"row": row, "name": unit.get("name") or row,
                     "labels": [[lb, unit_labels[lb]] for lb in order
                                if unit_labels.get(lb) not in (None, "")],
                     "questions": entries})
    return {"name": proj.get("name") or project.slug, "edition": proj.get("edition"),
            "unit_word": word, "unit_word_plural": words, "made": made, "questions": questions,
            "rows": rows}


# --- shared formatting ----------------------------------------------------------------------

def _label_name(label: str) -> str:
    return str(label).replace("_", " ").capitalize()


def long_date(value) -> str:
    """"2025-03-01" -> "1 March 2025"; anything else unchanged (as text)."""
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(value or ""))
    if not m or not 1 <= int(m.group(2)) <= 12:
        return "" if value is None else str(value)
    return f"{int(m.group(3))} {_MONTHS[int(m.group(2)) - 1]} {m.group(1)}"


def _benchmark(pack: dict) -> str:
    return f"{pack['name']}, {pack['edition']} edition" if pack.get("edition") else str(pack["name"])


def _title(pack: dict) -> str:
    rows = pack["rows"]
    if len(rows) == 1:
        return f"Evidence for {rows[0]['name']}"
    return f"Evidence for {len(rows)} {pack['unit_word_plural']}"


def _slug(text: str) -> str:
    return re.sub(r"[^a-z0-9.]+", "-", str(text).lower()).strip("-.")


def default_stem(rows: list[str] | None = None, label: str | None = None, all_rows: bool = False,
                 questions: list[str] | None = None) -> str:
    """The default file name without its extension: evidence-all, evidence-<label>-<value> or
    evidence-<row ids>, then -questions-<ids> when limited to questions. Long lists are cut down
    to the first item and a count."""
    if all_rows:
        what = "all"
    elif label is not None:
        name, _, value = str(label).partition("=")
        what = _slug(f"{name}-{value}")[:60].strip("-.") or "label"
    else:
        ids = [_slug(r) for r in rows or []]
        what = "-".join(ids)
        if len(what) > 50:
            what = f"{ids[0][:30]}-and-{len(ids) - 1}-more"
    stem = f"evidence-{what}"
    if questions:
        qs = "-".join(_slug(q) for q in questions)
        stem += f"-questions-{qs}" if len(qs) <= 30 else f"-{len(questions)}-questions"
    return stem


# --- the workbook ---------------------------------------------------------------------------

def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip().strip("'").strip()[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.casefold() in used:
        tail = f" ({n})"
        candidate, n = base[:31 - len(tail)].rstrip() + tail, n + 1
    used.add(candidate.casefold())
    return candidate


def _lines(entry: dict) -> list[list]:
    """The sheet lines for one question: one per quote, or one saying why there is none."""
    head = [entry["question"], entry["wording"], entry["answer"], entry["meaning"], entry["points"],
            entry["out_of"], entry["state"]]
    tail = ["; ".join(entry["flags"]) or None, "\n".join(entry["rules"]) or None]
    if not entry["evidence"]:
        return [head + [entry["note"]] + [None] * 7 + tail]
    return [head + [e["quote"], e["quote_en"], e["title"], e["publisher"], e["link"], e["page"],
                    e["date"], e["strength"]] + tail for e in entry["evidence"]]


def write_xlsx(pack: dict, path) -> Path:
    """Write the pack as a workbook, one sheet per row (its name, labels and the benchmark above a
    filterable table of COLUMNS). Text is always stored as text, never as a formula."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise HandbookError("Writing an Excel file needs the openpyxl package (pip install openpyxl). "
                            "Use --format html for a web page instead.") from e
    widths = {"Question": 10, "Wording": 40, "Answer": 10, "What it means": 30, "Points": 8,
              "Out of": 8, "State": 14, "Quote": 60, "Quote in English": 60, "Document title": 30,
              "Publisher": 20, "Link": 40, "Page": 7, "Document date": 13, "Source strength": 26,
              "Flags": 26, "Rules it rests on": 40}
    wrap = Alignment(wrap_text=True, vertical="top")
    bold = Font(bold=True)

    def put(ws, r, c, value):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        cell = ws.cell(row=r, column=c, value=value)
        if isinstance(value, str) and value.startswith("="):
            cell.data_type = "s"
        return cell

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for row in pack["rows"]:
        ws = wb.create_sheet(_sheet_name(row["name"], used))
        put(ws, 1, 1, row["name"]).font = Font(bold=True, size=14)
        put(ws, 2, 1, "; ".join(f"{_label_name(k)}: {v}" for k, v in row["labels"]) or None)
        put(ws, 3, 1, f"{_benchmark(pack)}. Prepared on {long_date(pack['made'])}.")
        head = 5
        for c, name in enumerate(COLUMNS, 1):
            put(ws, head, c, name).font = bold
            ws.column_dimensions[get_column_letter(c)].width = widths.get(name, 15)
        r = head
        for entry in row["questions"]:
            for line in _lines(entry):
                r += 1
                for c, value in enumerate(line, 1):
                    cell = put(ws, r, c, value)
                    cell.alignment = wrap
                    if COLUMNS[c - 1] == "Link" and isinstance(value, str) and \
                            value.lower().startswith(("http://", "https://")):
                        cell.hyperlink = value
        ws.freeze_panes = ws.cell(row=head + 1, column=1)
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(COLUMNS))}{max(r, head)}"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    return path


# --- the page -------------------------------------------------------------------------------

CSS = """
:root { color-scheme: light; }
body { margin: 0; background: #ffffff; color: #1d1d1f;
  font: 16px/1.55 -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; }
main { max-width: 46rem; margin: 0 auto; padding: 2rem 1rem 4rem; }
h1 { font-size: 1.8rem; line-height: 1.2; margin: 0 0 .4rem; }
h2 { font-size: 1.4rem; margin: 0 0 .2rem; padding-top: 1rem; border-top: 2px solid #1d1d1f; }
h3 { font-size: 1.05rem; margin: 0 0 .3rem; }
.lead { color: #55555a; margin: 0 0 1rem; }
.about p { margin: .4rem 0; color: #3a3a3f; }
nav ul { padding-left: 1.2rem; }
section.row { margin-top: 2.5rem; }
.labels { color: #55555a; margin: 0 0 1rem; }
article.q { margin: 0 0 1.25rem; padding: .9rem 1rem; border: 1px solid #d8d8dc; border-radius: 6px;
  break-inside: avoid; }
.qid { color: #55555a; font-weight: 600; margin-right: .15rem; }
.wording { margin: 0 0 .6rem; }
dl.answer { display: grid; grid-template-columns: max-content 1fr; gap: .15rem 1rem; margin: 0 0 .6rem; }
dl.answer dt { color: #55555a; }
dl.answer dd { margin: 0; }
.ev { margin: .7rem 0 0; padding-top: .6rem; border-top: 1px dashed #d8d8dc; }
blockquote { margin: 0 0 .35rem; padding: .1rem 0 .1rem .8rem; border-left: 3px solid #8a8a90; }
.tr, .src, .strength, .flags, .none { margin: .2rem 0; }
.tr span, .src span, .strength span, .flags span { color: #55555a; }
.strength.weak { color: #8a4b00; }
.none { font-style: italic; }
.rules { margin: .5rem 0 0; }
.rules p { margin: 0; color: #55555a; }
.rules ul { margin: .2rem 0 0; padding-left: 1.2rem; }
a { color: #0b57d0; overflow-wrap: anywhere; }
@media print { nav { display: none; } section.row + section.row { break-before: page; }
  a { color: inherit; } }
"""


def _e(value) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def _link(url) -> str:
    if isinstance(url, str) and url.lower().startswith(("http://", "https://")):
        return f'<a href="{_e(url)}">{_e(url)}</a>'
    return _e(url)


def _lang(code) -> str:
    return f' lang="{_e(code)}"' if isinstance(code, str) and re.fullmatch(r"[a-z]{2,3}(-[A-Za-z0-9]{2,8})?", code) else ""


def _points(entry: dict) -> str | None:
    if entry["points"] is None:
        return None
    return f"{entry['points']} of {entry['out_of']}" if entry["out_of"] is not None else str(entry["points"])


def _html_entry(entry: dict) -> list[str]:
    out = ['<article class="q">',
           f'<h3><span class="qid">{_e(entry["question"])}</span> {_e(entry["short"])}</h3>']
    if entry["wording"]:
        out.append(f'<p class="wording">{_e(entry["wording"])}</p>')
    facts = [("Answer", entry["answer"]), ("What it means", entry["meaning"]),
             ("Points", _points(entry)), ("State", entry["state"])]
    out.append('<dl class="answer">' + "".join(f"<dt>{_e(k)}</dt><dd>{_e(v)}</dd>" for k, v in facts
                                               if v not in (None, "")) + "</dl>")
    for ev in entry["evidence"]:
        out.append('<div class="ev">')
        out.append(f'<blockquote{_lang(ev["language"])}>{_e(ev["quote"])}</blockquote>')
        if ev["quote_en"]:
            out.append(f'<p class="tr" lang="en"><span>In English:</span> {_e(ev["quote_en"])}</p>')
        src = [_e(ev["title"])]
        if ev["publisher"]:
            src.append(_e(ev["publisher"]))
        if ev["date"]:
            src.append(f"Dated {_e(long_date(ev['date']))}")
        if ev["page"] is not None:
            src.append(f"Page {_e(ev['page'])}")
        out.append(f'<p class="src"><span>Source:</span> {". ".join(src)}.</p>')
        if ev["link"]:
            out.append(f'<p class="src"><span>Link:</span> {_link(ev["link"])}</p>')
        if ev["strength"]:
            kind = "weak" if ev["strength"] != OWN_WORDS else "own"
            out.append(f'<p class="strength {kind}"><span>Source strength:</span> {_e(ev["strength"])}</p>')
        out.append("</div>")
    if entry["note"]:
        out.append(f'<p class="none">{_e(entry["note"])}</p>')
    if entry["flags"]:
        out.append(f'<p class="flags"><span>Flags:</span> {_e("; ".join(entry["flags"]))}</p>')
    if entry["rules"]:
        out.append('<div class="rules"><p>Rules this answer rests on:</p><ul>'
                   + "".join(f"<li>{_e(r)}</li>" for r in entry["rules"]) + "</ul></div>")
    out.append("</article>")
    return out


def render_html(pack: dict) -> str:
    """One self-contained page: a title, fixed words on how to read it, a contents list when there
    is more than one row, then a section per row (id = the row id) with an entry per question."""
    title = _title(pack)
    out = ["<!doctype html>", '<html lang="en-GB">', "<head>", '<meta charset="utf-8">',
           '<meta name="viewport" content="width=device-width, initial-scale=1">',
           f"<title>{_e(title)}</title>", f"<style>{CSS}</style>", "</head>", "<body>", "<main>",
           f"<h1>{_e(title)}</h1>",
           f'<p class="lead">{_e(_benchmark(pack))}. Prepared on {_e(long_date(pack["made"]))}.</p>',
           '<div class="about">',
           "<p>Each answer below shows the words it rests on, copied exactly from the document "
           "named, with the page, the document's date and a link. Where a document is not in "
           "English, a translation follows the quote.</p>",
           "<p>Source strength says whether a document is in the organisation's own words or is a "
           "weaker source, such as a news article or a directory, and why. Where nothing was found "
           "for a question, this page says so.</p>",
           "</div>"]
    if len(pack["rows"]) > 1:
        out.append("<nav><h2>Contents</h2><ul>" + "".join(
            f'<li><a href="#{_e(r["row"])}">{_e(r["name"])}</a></li>' for r in pack["rows"]) + "</ul></nav>")
    for row in pack["rows"]:
        out.append(f'<section id="{_e(row["row"])}" class="row">')
        out.append(f"<h2>{_e(row['name'])}</h2>")
        if row["labels"]:
            out.append('<p class="labels">' + _e(". ".join(f"{_label_name(k)}: {v}" for k, v in row["labels"]))
                       + ".</p>")
        if not row["questions"]:
            out.append('<p class="none">None of the chosen questions apply here.</p>')
        for entry in row["questions"]:
            out.extend(_html_entry(entry))
        out.append("</section>")
    out += ["</main>", "</body>", "</html>", ""]
    return "\n".join(out)


# --- the export -----------------------------------------------------------------------------

def export(project: Project, rows: list[str] | None = None, label: str | None = None,
           all_rows: bool = False, questions: list[str] | None = None, fmt: str | None = None,
           out=None, now: datetime | None = None) -> dict:
    """Choose the rows and questions, build the pack and write it. `fmt` defaults to the ending of
    `out` (.html or .xlsx), else xlsx. `out` defaults to `<project>/exports/<default_stem>.<fmt>`;
    an existing folder gets the default name inside it. Returns {"path", "format", "rows",
    "questions", "with_evidence"} (counts of rows, question entries, entries with a quote)."""
    units = choose_rows(project, rows, label, all_rows)
    qids = choose_questions(project, questions)
    target = Path(out).expanduser() if out else None
    if fmt is None:
        suffix = target.suffix.lower() if target is not None else ""
        fmt = "html" if suffix in (".html", ".htm") else "xlsx"
    if fmt not in FORMATS:
        raise HandbookError(f"The format must be one of {', '.join(FORMATS)}, not {fmt!r}.")
    stem = default_stem([u.get("id") for u in units] if rows is not None else None,
                        label, all_rows, qids)
    if target is None:
        target = project.path / "exports" / f"{stem}.{fmt}"
    elif target.is_dir():
        target = target / f"{stem}.{fmt}"
    pack = build_pack(project, units, qids, now)
    if fmt == "xlsx":
        write_xlsx(pack, target)
    else:
        target.parent.mkdir(parents=True, exist_ok=True)
        tmp = target.with_name(target.name + ".tmp")
        try:
            tmp.write_text(render_html(pack), encoding="utf-8")
            os.replace(tmp, target)
        except BaseException:
            tmp.unlink(missing_ok=True)
            raise
    entries = [e for r in pack["rows"] for e in r["questions"]]
    return {"path": target, "format": fmt, "rows": len(pack["rows"]), "questions": len(entries),
            "with_evidence": sum(1 for e in entries if e["evidence"])}
