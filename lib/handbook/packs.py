"""Packs: the Excel file a party gets with its results, and the reply sheet it fills in.

`make_packs` writes one pack per unit, or one per group of units (a label's value), into
`<project>/packs/<YYYY-MM-DD>/`. A pack has

- "Start here": what this is, the headline score, how to reply, who to send it to;
- "Your scores" (a group pack has one tab per unit, named after the unit): one row per question
  with the score, what it means, how to score higher, why we gave it, and the evidence we used.
  Those columns are locked. The yellow reply columns are unlocked: "Do you agree?", why or what
  is new, a link, a page or section;
- a very hidden sheet `_pack` holding the unit ids, our values, a fingerprint of each answer and
  one of each sheet's locked columns, so a reply can be matched to its unit, a reply to an older
  pack can be flagged ("we have changed this answer since"), and edits to our columns noticed.

Nothing in the text is specific to one benchmark: names, words, sections, levels and the score
sheet all come from the project. Making a pack records "pack made" and moves the unit to ready to
send (feedback.record_packs).
"""
from __future__ import annotations
import hashlib
import json
import os
from datetime import date, datetime
from pathlib import Path

from . import feedback
from .folder import HandbookError, Project
from .scoresheet import round_half_up, row_totals

PACK_VERSION = 2
PACK_SHEET = "_pack"
START_SHEET = "Start here"
SCORES_SHEET = "Your scores"
CHUNK = 30000                     # an Excel cell holds at most 32,767 characters

STRINGS = {
    "en": {
        "title": "{title}: your results",
        "intro_unit": "These are the results we've worked out for {name}. Before they're final, we'd like you to check them.",
        "intro_group": "These are the results we've worked out for {count} {words} under {name}. "
                       "Before they're final, we'd like you to check them.",
        "overall": "Overall score",
        "headline": "Score",
        "tier": "tier",
        "how": "How to reply",
        "step1_unit": "1. Open the ‘{sheet}’ tab. Each row is one question, with your score and the evidence we used.",
        "step1_group": "1. Open the tab for each {word} (they are listed above). Each row is one question, "
                       "with the score and the evidence we used.",
        "step2_public": "2. In the yellow columns, say whether you agree. If you don't, tell us why and give a "
                        "link to a document anyone can open.",
        "step2_private": "2. In the yellow columns, say whether you agree. If you don't, tell us why and give a "
                         "link, or attach the document to your email. Your own documents, such as internal reports, "
                         "are fine; we will use them only to check this answer.",
        "step3": "3. Save this file and email it back to {sender}{by}. Please keep ‘{keep}’ in the file name, "
                 "so your reply reaches the right place.",
        "by": " by {date}",
        "blank": "You only need to fill in rows you want to comment on. Blank rows count as no comment.",
        "tabs": "Tabs in this file",
        "no_evidence": "We found no evidence on this.",
        "top": "You have the top score here.",
        "score_of": "{got} of {top}",
        "columns": ["Q", "Topic", "What we asked", "Your score", "What this score means", "To score higher",
                    "Why we gave this score", "Evidence we used"],
        "reply_columns": ["Do you agree?", "If not, why? Or what's new?", "Link to your evidence", "Page or section"],
        "choices": ["Agree", "Disagree", "Agree, with new information"],
        "page": ", p.{page}",
    },
}

INK, MUTED, GREEN, PALE_GREEN, YELLOW = "1F2933", "616E7C", "1E7F5C", "E8F5EE", "FFF6D6"


def strings(language: str | None) -> dict:
    """The pack's words in `language`; English when there is no translation (yet)."""
    return STRINGS.get((language or "en").lower().split("-")[0], STRINGS["en"])


def _s(value) -> str:
    return "" if value is None else str(value)


def sheet_fingerprint(rows: list[list]) -> str:
    """A fingerprint of a sheet's locked columns, as the rows' text (so a reader can tell whether
    anyone edited our side of the file)."""
    return hashlib.sha1(json.dumps([[_s(v) for v in r] for r in rows], sort_keys=True,
                                   ensure_ascii=False).encode("utf-8")).hexdigest()[:12]


# --- what goes in a pack --------------------------------------------------------------------------

def _doc_meta(project: Project, doc: str, cache: dict) -> dict:
    if doc not in cache:
        cache[doc] = project.documents().get(doc) or {"title": doc, "url": ""}
    return cache[doc]


def _evidence_text(project: Project, cell: dict, st: dict, cache: dict) -> str:
    lines = []
    for e in cell.get("evidence") or []:
        meta = _doc_meta(project, e.get("doc"), cache)
        quote = e.get("quote_en") or e.get("quote") or ""
        page = st["page"].format(page=e["page"]) if e.get("page") else ""
        lines.append(f"“{quote}”\n— {meta.get('title') or ''}{page}\n{meta.get('url') or ''}".rstrip())
    return "\n\n".join(lines) or st["no_evidence"]


def _next_level(question: dict, points, st: dict) -> str:
    levels = (question.get("answer") or {}).get("levels") or []
    above = [lv for lv in levels if isinstance(lv.get("points"), (int, float)) and lv["points"] > (points or 0)]
    return min(above, key=lambda lv: lv["points"]).get("meaning") or "" if above else st["top"]


def _top(question: dict):
    return max((lv.get("points") or 0 for lv in (question.get("answer") or {}).get("levels") or []), default=0)


def _num(x) -> str:
    return str(int(x)) if isinstance(x, float) and x.is_integer() else str(x)


def unit_data(project: Project, unit: dict, st: dict, settings: dict, state_rec: dict | None = None) -> dict | None:
    """Everything a pack says about one unit, or None when it has no answers to show yet."""
    uid = unit["id"]
    cells = project.answers(uid).get("cells") or {}
    questions = [q for q in project.questions_for(unit)
                 if isinstance(cells.get(q["id"]), dict) and cells[q["id"]].get("state") in ("answered", "no_evidence")]
    if not questions:
        return None
    sheet = project.read("scoresheet") or {}
    plain = (state_rec or {}).get("plain_reasons") or {} if settings.get("reasons") == "plain" else {}
    plain_for = (state_rec or {}).get("plain_for") or {}      # the answer each plain reason was written for
    cache: dict = {}
    rows, fixed, our_values, fps = [], [], {}, {}
    for q in questions:
        cell = cells[q["id"]]
        value = cell.get("value") if cell.get("state") == "answered" else q.get("no_evidence_value")
        level = next((lv for lv in (q.get("answer") or {}).get("levels") or [] if lv.get("value") == value), None)
        got = cell.get("points") if cell.get("points") is not None else (level or {}).get("points", 0)
        fp = feedback.answer_fingerprint(cell)
        # A plain reason is used only while the answer is the one it was written for.
        reason = plain.get(q["id"]) if plain_for.get(q["id"], fp) == fp else None
        vals = [q["id"], q.get("short") or q["id"], q.get("wording") or "",
                st["score_of"].format(got=_num(got), top=_num(_top(q))),
                (level or {}).get("meaning") or "", _next_level(q, got, st),
                reason or cell.get("note") or "", _evidence_text(project, cell, st, cache)]
        fixed.append(vals)
        our_values[q["id"]] = cell.get("value")
        fps[q["id"]] = fp
        rows.append(vals)
    totals = row_totals(project, unit, cells)
    in_pack = {q.get("section") for q in questions}
    sections, overall, band = [], None, None
    if totals:
        overall, band = totals["overall"], (totals.get("bands") or {}).get("overall")
        for s in sheet.get("sections") or []:
            entry = totals["sections"].get(s["id"]) or {}
            if s["id"] not in in_pack or entry.get("percent") is None:
                continue
            if s.get("percent", True):
                sections.append((s["id"], s.get("name") or s["id"], f"{_num(entry['percent'])}%"))
            else:
                sections.append((s["id"], s.get("name") or s["id"], _num(round_half_up(entry["points"], s.get("rounding")))))
    ov = sheet.get("overall") or {}
    overall_is_percent = ov.get("method") == "weighted_mean_of_section_percents" or bool(ov.get("percent"))
    return {"unit": uid, "name": unit.get("name") or uid, "rows": rows, "questions": [q["id"] for q in questions],
            "our_values": our_values, "answers": fps, "fingerprint": sheet_fingerprint(fixed),
            "overall": overall, "overall_is_percent": overall_is_percent, "band": band, "sections": sections}


def headline_text(data: dict, settings: dict, st: dict) -> tuple[str, str] | None:
    """(label, value) for the pack's headline line: the overall score by default, or one section."""
    mode = settings.get("headline", "overall")
    if mode == "none":
        return None
    if mode == "overall":
        if data["overall"] is None:
            return None
        value = f"{_num(data['overall'])}{'%' if data['overall_is_percent'] else ''}"
        return st["overall"], value + (f", {data['band']}" if data["band"] else "")
    for sid, name, value in data["sections"]:
        if sid == mode:
            return name, value
    return None


# --- the workbook --------------------------------------------------------------------------------

def _tab_name(name: str, used: set) -> str:
    base = "".join("_" if c in '[]:*?/\\' else c for c in str(name)).strip("' ")[:31] or "Unit"
    out, n = base, 2
    while out.lower() in used or out in (START_SHEET, PACK_SHEET):
        suffix = f" ({n})"
        out, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(out.lower())
    return out


def _scores_sheet(wb, title: str, data: dict, st: dict):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    line = Side(style="thin", color="D5DBE1")
    ws = wb.create_sheet(title)
    head = st["columns"] + st["reply_columns"]
    widths = [6, 22, 38, 10, 34, 34, 46, 60, 16, 44, 34, 16]
    for i, (h, w) in enumerate(zip(head, widths), 1):
        c = ws.cell(1, i, h)
        reply = i > len(st["columns"])
        c.font = Font(bold=True, color=INK if reply else "FFFFFF")
        c.fill = PatternFill("solid", fgColor="F2C94C" if reply else GREEN)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    choices = DataValidation(type="list", formula1='"' + ",".join(st["choices"]) + '"', allow_blank=True)
    ws.add_data_validation(choices)
    for n, vals in enumerate(data["rows"], 2):
        for i, v in enumerate(vals, 1):
            c = ws.cell(n, i, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=line)
            c.font = Font(color=INK, bold=(i == 4))
        for i in range(len(vals) + 1, len(head) + 1):
            c = ws.cell(n, i)
            c.fill = PatternFill("solid", fgColor=YELLOW)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(bottom=line, left=line)
            c.protection = Protection(locked=False)
        choices.add(ws.cell(n, len(vals) + 1))
        ws.row_dimensions[n].height = 150
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(head)).column_letter}{len(data['rows']) + 1}"
    return ws


def _start_sheet(wb, project: Project, items: list[dict], group: dict | None, settings: dict, st: dict,
                 reply_by: date | None, sender: str, tabs: dict):
    from openpyxl.styles import Alignment, Font, PatternFill
    proj = project.read("project")
    word = proj.get("unit_word") or "unit"
    words = proj.get("unit_word_plural") or "units"
    name = proj.get("name") or "Benchmark"
    title = name if str(proj.get("edition") or "") in name else f"{name} {proj.get('edition') or ''}".strip()
    ws = wb.active
    ws.title = START_SHEET
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 70
    r = 2

    def merged(text, font, height=None):
        nonlocal r
        c = ws.cell(r, 2, text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        if height:
            ws.row_dimensions[r].height = height
        r += 1

    def kv(label, value, bold=False, fill=None):
        nonlocal r
        a, b = ws.cell(r, 2, label), ws.cell(r, 3, value)
        a.font = Font(color=MUTED)
        b.font = Font(bold=bold, color=INK, size=12 if bold else 11)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            for c in (a, b):
                c.fill = PatternFill("solid", fgColor=fill)
        r += 1

    shown = group["value"] if group else items[0]["name"]
    merged(st["title"].format(title=title), Font(size=18, bold=True, color=INK))
    merged(shown, Font(size=13, color=MUTED))
    r += 1
    intro = (st["intro_group"].format(count=len(items), words=words, name=shown) if group
             else st["intro_unit"].format(name=shown))
    merged(intro, Font(size=11, color=INK), 32)
    r += 1
    if group:
        merged(st["tabs"], Font(size=13, bold=True, color=INK))
        for d in items:
            h = headline_text(d, settings, st)
            kv(d["name"], (f"{h[0]}: {h[1]}" if h else "") + f"  (tab ‘{tabs[d['unit']]}’)")
    else:
        h = headline_text(items[0], settings, st)
        if h:
            kv(h[0], h[1], bold=True, fill=PALE_GREEN)
            if settings.get("headline", "overall") == "overall":
                for _, sname, value in items[0]["sections"]:
                    kv(sname, value)
    r += 1
    merged(st["how"], Font(size=13, bold=True, color=INK))
    by = st["by"].format(date=f"{reply_by.day} {reply_by:%B %Y}") if reply_by else ""
    steps = [st["step1_group"].format(word=word) if group else st["step1_unit"].format(sheet=SCORES_SHEET),
             st["step2_private"] if settings.get("private_evidence") is True else st["step2_public"],
             st["step3"].format(sender=sender, by=by, keep=shown)]
    for s in steps:
        merged(s, Font(color=INK), 32)
    r += 1
    merged(st["blank"], Font(italic=True, color=MUTED))
    return title


def _pack_sheet(wb, pack: dict):
    ws = wb.create_sheet(PACK_SHEET)
    text = json.dumps(pack, ensure_ascii=False, separators=(",", ":"))
    for i in range(0, len(text), CHUNK):
        ws.cell(1 + i // CHUNK, 1, text[i:i + CHUNK])
    ws.sheet_state = "veryHidden"
    ws.protection.sheet = True


def build_workbook(project: Project, items: list[dict], group: dict | None, settings: dict, reply_by: date | None,
                   sender: str, st: dict, today: date):
    """The workbook for one unit (`group` None) or a group of units, from `unit_data` items."""
    from openpyxl import Workbook
    proj = project.read("project")
    wb = Workbook()
    used: set = set()
    tabs = {}
    for d in items:
        tabs[d["unit"]] = _tab_name(SCORES_SHEET if not group else d["name"], used)
    title = _start_sheet(wb, project, items, group, settings, st, reply_by, sender, tabs)
    for d in items:
        _scores_sheet(wb, tabs[d["unit"]], d, st)
    pack = {"pack_version": PACK_VERSION, "benchmark": proj.get("slug"), "edition": proj.get("edition"),
            "made": today.isoformat(), "reply_by": reply_by.isoformat() if reply_by else None,
            "group": group, "choices": st["choices"],
            "units": [{"unit": d["unit"], "unit_name": d["name"], "sheet": tabs[d["unit"]],
                       "questions": d["questions"], "our_values": d["our_values"], "answers": d["answers"],
                       "fingerprint": d["fingerprint"]} for d in items]}
    _pack_sheet(wb, pack)
    return wb, title


# --- making packs --------------------------------------------------------------------------------

def _parse_day(value, what: str) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        raise HandbookError(f"{what} must be a date like 2026-11-13, not {value!r}.") from None


def choose_units(project: Project, scope: str, units=None) -> list[dict]:
    """The units a make covers: `scope` "units" (the ids or names given), "ready" (those at ready
    to send) or "all"."""
    every = project.read("units").get("units") or []
    if scope == "all":
        return list(every)
    if scope == "units":
        if not units:
            raise HandbookError("Say which units to make packs for.")
        ids = [feedback.unit_id(project, u) for u in ([units] if isinstance(units, str) else units)]
        return [u for u in every if u["id"] in ids]
    if scope == "ready":
        state = feedback.read_state(project)
        return [u for u in every
                if feedback.stage_of(feedback.unit_rec(state, u["id"]), len(feedback.unit_unchecked(project, u["id"], u)))
                == "ready_to_send"]
    raise HandbookError("Packs can be made for units, ready or all.")


def make_packs(project: Project, scope: str = "ready", units=None, group_by: str | None = None,
               reply_by=None, sender: str | None = None, now: datetime | None = None) -> dict:
    """Write packs for the chosen units and record them. `group_by` is a label of the project
    ("brand") to make one pack per value of it, with a tab per unit; "none" makes one per unit;
    None takes the grouping setting. Units whose answers still need checking are made anyway and
    listed in `warnings` ({unit: [questions]}); units with no answers yet are `skipped`. Returns
    {"folder", "packs": [{"file", "name", "units", "group"}], "warnings", "skipped", "plain_missing"} (the
    last: units made with analyst notes although plain reasons are asked for, because none is written yet)."""
    settings = feedback.read_settings(project)
    proj = project.read("project")
    mode = group_by if group_by is not None else settings["grouping"]
    if mode is None or str(mode).strip().lower() in ("", "none", "unit", "units"):
        mode = None
    elif mode not in (proj.get("labels") or []):
        raise HandbookError(f"Packs can be grouped by one of the labels ({', '.join(proj.get('labels') or []) or 'none set'}), "
                            f"or one per {feedback.unit_word(project)}. Not {mode!r}.")
    reply_by = _parse_day(reply_by if reply_by is not None else settings.get("reply_by"), "The reply-by date")
    sender = (sender or settings["sender"]).strip()
    st = strings(settings.get("language") or proj.get("language"))
    state = feedback.read_state(project)
    chosen = choose_units(project, scope, units)
    if not chosen:
        raise HandbookError("There are no units to make packs for.")
    today = feedback._now(now).date()

    items, skipped, warnings, plain_missing = [], [], {}, []
    for unit in chosen:
        data = unit_data(project, unit, st, settings, state["units"].get(unit["id"]))
        if data is None:
            skipped.append({"unit": unit["id"], "why": "It has no answers to show yet."})
            continue
        items.append((unit, data))
        if settings.get("reasons") == "plain" and not (state["units"].get(unit["id"]) or {}).get("plain_reasons"):
            plain_missing.append(unit["id"])      # plain wording asked for, not written for this one yet
        todo = feedback.unit_unchecked(project, unit["id"], unit)
        if todo:
            warnings[unit["id"]] = todo
    if not items:
        raise HandbookError("None of those units has answers to show yet, so no packs were made.")

    groups: list[tuple[dict | None, list]] = []
    if mode:
        by_value: dict[str, list] = {}
        for unit, data in items:
            value = (unit.get("labels") or {}).get(mode)
            if value:
                by_value.setdefault(value, []).append((unit, data))
            else:
                groups.append((None, [(unit, data)]))
        for value, members in by_value.items():
            groups.append(({"label": mode, "value": value} if len(members) > 1 else None, members))
    else:
        groups = [(None, [it]) for it in items]

    folder = project.path / "packs" / today.isoformat()
    folder.mkdir(parents=True, exist_ok=True)
    made, packs = [], []
    for group, members in groups:
        wb, title = build_workbook(project, [d for _, d in members], group, settings, reply_by, sender, st, today)
        label = group["value"] if group else members[0][1]["name"]
        path = folder / f"{feedback.safe_name(title + ' - ' + label)}.xlsx"
        tmp = path.with_name(path.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, path)
        rel = path.relative_to(project.path).as_posix()
        ids = [u["id"] for u, _ in members]
        made.append({"units": ids, "file": rel, "group": group})
        packs.append({"file": rel, "name": label, "units": ids, "group": group})
    feedback.record_packs(project, made, now)
    return {"folder": folder.relative_to(project.path).as_posix(), "packs": packs, "warnings": warnings,
            "skipped": skipped, "plain_missing": plain_missing}
